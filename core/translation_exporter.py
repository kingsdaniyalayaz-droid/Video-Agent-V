"""Standalone in-memory exports for completed Roman Urdu translations.

This module intentionally has no dependency on Streamlit, the translation
pipeline, Mistral, Whisper, or any other application service.
"""

from __future__ import annotations

from datetime import datetime
from html import escape
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


_TIMESTAMP_PATTERN = re.compile(
    r"(?:\[\s*\d{1,2}:\d{2}(?::\d{2})?(?:\.\d+)?\s*\]|"
    r"\b\d{1,2}:\d{2}:\d{2}(?:\.\d+)?\b)"
)


def _export_date() -> str:
    """Return a human-readable export timestamp."""
    return datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")


def _parse_timestamp_blocks(transcript: str) -> list[dict[str, str]]:
    """Parse timestamp-owned transcript blocks while preserving source order."""
    if not transcript or not transcript.strip():
        return []
    matches = list(_TIMESTAMP_PATTERN.finditer(transcript))
    if not matches:
        return []

    blocks: list[dict[str, str]] = []
    prefix = transcript[:matches[0].start()]
    if prefix.strip():
        blocks.append({"timestamp": "", "text": prefix.strip()})

    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(transcript)
        blocks.append({
            "timestamp": match.group(0),
            "text": transcript[match.end():end].strip(),
        })
    return blocks


def _plain_record(transcript: str) -> list[dict[str, str]]:
    """Represent an untimestamped transcript as one complete record."""
    return [{"timestamp": "", "text": transcript or ""}]


def _aligned_records(original_transcript: str, roman_urdu_transcript: str) -> list[dict[str, str]]:
    """Align original and translated records using exact source timestamps."""
    original_blocks = _parse_timestamp_blocks(original_transcript)
    translation_blocks = _parse_timestamp_blocks(roman_urdu_transcript)

    if not original_blocks:
        return [{
            "timestamp": "",
            "original": original_transcript or "",
            "roman_urdu": roman_urdu_transcript or "",
        }]

    source_timestamps = [block["timestamp"] for block in original_blocks]
    translated_timestamps = [block["timestamp"] for block in translation_blocks]
    if translated_timestamps == source_timestamps and len(translation_blocks) == len(original_blocks):
        translated_text = [block["text"] for block in translation_blocks]
    else:
        plain_translation = [
            line.strip()
            for line in (roman_urdu_transcript or "").splitlines()
            if line.strip()
        ]
        if len(plain_translation) == len(original_blocks):
            translated_text = plain_translation
        else:
            # Preserve the complete supplied translation rather than dropping
            # text when an external caller supplies an unusual segmentation.
            translated_text = [roman_urdu_transcript or ""] + [""] * (len(original_blocks) - 1)

    return [
        {
            "timestamp": block["timestamp"],
            "original": block["text"],
            "roman_urdu": translated_text[index],
        }
        for index, block in enumerate(original_blocks)
    ]


def build_translation_txt(
    original_transcript: str,
    roman_urdu_transcript: str,
    title: str = "Roman Urdu Translation",
) -> bytes:
    """Build a UTF-8 TXT export containing aligned original and translation text."""
    records = _aligned_records(original_transcript, roman_urdu_transcript)
    lines = [title, f"Export Date: {_export_date()}", ""]
    for record in records:
        if record["timestamp"]:
            lines.append(record["timestamp"])
        lines.extend([
            "ORIGINAL:",
            record["original"],
            "ROMAN URDU:",
            record["roman_urdu"],
            "",
        ])
    return "\n".join(lines).encode("utf-8")


def build_translation_excel(
    original_transcript: str,
    roman_urdu_transcript: str,
) -> bytes:
    """Build an in-memory XLSX export with aligned rows and metadata."""
    records = _aligned_records(original_transcript, roman_urdu_transcript)
    export_date = _export_date()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Roman Urdu Translation"
    headers = ["Sr. No.", "Timestamp", "Original Transcript", "Roman Urdu Translation"]
    sheet.append(headers)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index, record in enumerate(records, start=1):
        sheet.append([index, record["timestamp"], record["original"], record["roman_urdu"]])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
    for column, width in {"A": 10, "B": 18, "C": 60, "D": 60}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, max_col=4):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    metadata = workbook.create_sheet("Metadata")
    metadata.append(["Field", "Value"])
    metadata.append(["Export Date", export_date])
    metadata.append(["Total Segments", len(records)])
    metadata.append(["Original Characters", len(original_transcript or "")])
    metadata.append(["Roman Urdu Characters", len(roman_urdu_transcript or "")])
    for cell in metadata[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    metadata.column_dimensions["A"].width = 24
    metadata.column_dimensions["B"].width = 30
    metadata.freeze_panes = "A2"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _paragraph(text: str, style: ParagraphStyle) -> Paragraph:
    """Escape transcript text before placing it in a ReportLab Paragraph."""
    return Paragraph(escape(text or "").replace("\n", "<br/>"), style)


def build_translation_pdf(
    original_transcript: str,
    roman_urdu_transcript: str,
    title: str = "Roman Urdu Translation Report",
) -> bytes:
    """Build a readable multi-page in-memory PDF export."""
    records = _aligned_records(original_transcript, roman_urdu_transcript)
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=title,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ExportTitle", parent=styles["Title"], alignment=TA_CENTER, spaceAfter=8)
    meta_style = ParagraphStyle("ExportMeta", parent=styles["Normal"], alignment=TA_CENTER, textColor=colors.HexColor("#555555"), spaceAfter=12)
    heading_style = ParagraphStyle("ExportHeading", parent=styles["Heading2"], spaceBefore=8, spaceAfter=5)
    label_style = ParagraphStyle("ExportLabel", parent=styles["Heading3"], spaceBefore=5, spaceAfter=3)
    body_style = ParagraphStyle("ExportBody", parent=styles["BodyText"], leading=13, spaceAfter=4)
    timestamp_style = ParagraphStyle("ExportTimestamp", parent=body_style, textColor=colors.HexColor("#1F4E78"), fontName="Helvetica-Bold", spaceAfter=4)

    # Transcript text is emitted as independent flowables. This avoids putting
    # a potentially huge original or translated block inside a table row.
    story: list[Any] = [
        _paragraph(title, title_style),
        _paragraph(f"Export date: {_export_date()}", meta_style),
        _paragraph(f"Total segments: {len(records)}", meta_style),
    ]
    for record in records:
        if record["timestamp"]:
            story.append(_paragraph(record["timestamp"], timestamp_style))
        story.append(_paragraph("Original Transcript", label_style))
        story.append(_paragraph(record["original"], body_style))
        story.append(Spacer(1, 4))
        story.append(_paragraph("Roman Urdu Translation", label_style))
        story.append(_paragraph(record["roman_urdu"], body_style))
        story.append(Spacer(1, 10))
    document.build(story)
    return output.getvalue()


__all__ = [
    "build_translation_txt",
    "build_translation_excel",
    "build_translation_pdf",
]
